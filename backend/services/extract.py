import pypandoc
import os
from docx import Document
import numpy as np
from pathlib import Path


texte = r"""# Compte Rendu de TP

## Objectif
L'objectif de ce TP est de maitriser la loi d'Ohm aux bornes d'une resistance

## Formation requise
La relation est $$R = \frac{U}{I}$$
"""

#pypandoc.convert_text(texte, "docx", format="md", outputfile="mesures.docx")

folder = Path(__file__).resolve().parent
doc = Document(folder / "mesures.docx")

def nombre(texte):
    try:
        return float(texte.strip().replace(",", "."))
    except ValueError:
        return texte


def extraction(doc):
    
    result = []
    for i, tableau in enumerate(doc.tables):
        tab = {}
        for line in tableau.rows:
            nom = line.cells[0].text.strip()
            valeur = [nombre(cellule.text) for cellule in line.cells[1:]]
            tab[nom] = valeur
        result.append(tab)
    return result

extraction_tab = extraction(doc)

if __name__=="__main__":
    doc = Document(folder / "mesures.docx")
